from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
LOG_DIR = ROOT / "logs" / "assistive_fetch"
OUTPUT = ROOT / "assistive_fetch_interview_presentation.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(20, 34, 54)
NAVY_2 = RGBColor(32, 51, 75)
TEAL = RGBColor(25, 145, 145)
TEAL_LIGHT = RGBColor(221, 242, 240)
ORANGE = RGBColor(230, 132, 45)
ORANGE_LIGHT = RGBColor(251, 235, 218)
RED = RGBColor(194, 67, 67)
GREEN = RGBColor(55, 142, 94)
WHITE = RGBColor(255, 255, 255)
OFF_WHITE = RGBColor(247, 249, 251)
LIGHT_GRAY = RGBColor(225, 231, 237)
MID_GRAY = RGBColor(111, 125, 140)
DARK = RGBColor(31, 41, 51)

FONT_CN = "Microsoft YaHei"
FONT_MONO = "Consolas"


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font=FONT_CN,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, color=DARK, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT_CN
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(9)
        if bullet:
            p.text = f"• {line}"
    return box


def add_rect(slide, x, y, w, h, fill=WHITE, line=LIGHT_GRAY, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_card(slide, title, body, x, y, w, h, accent=TEAL):
    add_rect(slide, x, y, w, h, WHITE, LIGHT_GRAY, True)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, x + 0.22, y + 0.14, w - 0.35, 0.38, 18, NAVY, True)
    add_text(slide, body, x + 0.22, y + 0.58, w - 0.35, h - 0.7, 13.5, DARK)


def add_title(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, section.upper(), 0.55, 0.18, 2.6, 0.25, 9, TEAL, True)
    add_text(slide, title, 0.55, 0.45, 12.1, 0.55, 25, NAVY, True)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.07), Inches(1.0), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    if subtitle:
        add_text(slide, subtitle, 1.75, 1.00, 10.7, 0.32, 11, MID_GRAY)


def add_footer(slide, page, source=None):
    if source:
        add_text(slide, source, 0.55, 7.14, 10.9, 0.18, 7.5, MID_GRAY)
    add_text(slide, f"{page:02d}", 12.25, 7.08, 0.5, 0.22, 9, MID_GRAY, True, PP_ALIGN.RIGHT)


def add_connector(slide, x1, y1, x2, y2, color=MID_GRAY, width=2, arrow=True):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_process_box(slide, title, subtitle, x, y, w, h, fill, accent):
    add_rect(slide, x, y, w, h, fill, accent, True)
    add_text(slide, title, x + 0.12, y + 0.14, w - 0.24, 0.32, 16, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, subtitle, x + 0.12, y + 0.52, w - 0.24, h - 0.62, 10.5, DARK, False, PP_ALIGN.CENTER)


def add_native_chart(slide, categories, series, x, y, w, h, title, max_value=None, percent=False):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].font.name = FONT_CN
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = LIGHT_GRAY
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(9)
    if max_value is not None:
        chart.value_axis.maximum_scale = max_value
    if percent:
        chart.value_axis.tick_labels.number_format = "0%"
    palette = [TEAL, ORANGE, NAVY_2, GREEN]
    for i, plot_series in enumerate(chart.series):
        plot_series.format.fill.solid()
        plot_series.format.fill.fore_color.rgb = palette[i % len(palette)]
        plot_series.format.line.color.rgb = palette[i % len(palette)]
    return chart


def add_metric(slide, value, label, x, y, w, accent=TEAL, note=None):
    add_text(slide, value, x, y, w, 0.55, 28, accent, True, PP_ALIGN.CENTER)
    add_text(slide, label, x, y + 0.56, w, 0.35, 11, NAVY, True, PP_ALIGN.CENTER)
    if note:
        add_text(slide, note, x, y + 0.94, w, 0.32, 8.5, MID_GRAY, False, PP_ALIGN.CENTER)


def set_background(slide, color=OFF_WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1. Title
    slide = prs.slides.add_slide(blank)
    set_background(slide, NAVY)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.14), SLIDE_H
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_text(slide, "基于共享控制与深度强化学习的\nFetchPush 辅助操作", 0.85, 1.15, 11.4, 1.6, 31, WHITE, True)
    add_text(slide, "运动控制算法岗位 · 项目技术汇报", 0.88, 3.02, 7.0, 0.5, 17, RGBColor(183, 219, 218), True)
    add_text(
        slide,
        "规则控制器提供人类意图，SAC 学习受限幅值的辅助动作；\n通过状态机、奖励塑形和定量评估实现可靠共享控制。",
        0.88,
        3.72,
        8.9,
        1.05,
        16,
        WHITE,
    )
    add_rect(slide, 9.8, 3.35, 2.55, 1.7, NAVY_2, RGBColor(65, 91, 117), True)
    add_text(slide, "84%", 10.05, 3.61, 2.05, 0.6, 30, RGBColor(99, 215, 203), True, PP_ALIGN.CENTER)
    add_text(slide, "最佳辅助成功率", 10.05, 4.28, 2.05, 0.36, 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "候选人：________    日期：2026", 0.88, 6.72, 5.7, 0.32, 11, RGBColor(181, 192, 204))

    # 2. Problem
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "项目问题：让机器人“辅助”，而不是“接管”", "FetchPush-v4 · 连续控制 · 人机共享控制", "01 / Problem")
    add_card(
        slide,
        "任务",
        "机械臂末端推动立方体到随机目标位置；每回合 50 steps，成功阈值为目标距离 < 0.05 m。",
        0.55,
        1.45,
        3.75,
        2.0,
        TEAL,
    )
    add_card(
        slide,
        "输入",
        "启发式 human intent controller 给出基础动作；RL agent 只学习幅值受限的 assist action。",
        4.55,
        1.45,
        3.75,
        2.0,
        ORANGE,
    )
    add_card(
        slide,
        "目标",
        "提高成功率、缩短最终距离，同时控制辅助努力和动作变化率，保持人类主导。",
        8.55,
        1.45,
        3.75,
        2.0,
        GREEN,
    )
    add_text(slide, "三个核心挑战", 0.65, 3.88, 3.0, 0.38, 19, NAVY, True)
    challenges = [
        ("接触前规划", "末端必须先到物体后方并下降到合适高度，避免从上方/侧面撞击。"),
        ("有限时域", "50 steps 内完成 lift → approach → descend → push，前置阶段过慢会挤压推物时间。"),
        ("学习稳定性", "稀疏成功、随机初始状态、动作高频抖动会导致价值估计和控制品质不稳定。"),
    ]
    for i, (t, b) in enumerate(challenges):
        x = 0.65 + i * 4.05
        add_text(slide, f"{i + 1}", x, 4.48, 0.48, 0.48, 18, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(4.45), Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = TEAL
        circle.line.fill.background()
        circle.text_frame.text = str(i + 1)
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        circle.text_frame.paragraphs[0].runs[0].font.name = FONT_CN
        circle.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        circle.text_frame.paragraphs[0].runs[0].font.bold = True
        circle.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
        add_text(slide, t, x + 0.65, 4.43, 3.05, 0.35, 15, NAVY, True)
        add_text(slide, b, x, 5.08, 3.7, 1.15, 12.5, DARK)
    add_footer(slide, 2, "Source: assistive_fetch/fetch_push_wrapper.py · Gymnasium Robotics FetchPush-v4")

    # 3. Architecture
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "共享控制闭环：规则意图 + 学习辅助 + 物理反馈", "Agent 不直接替代 human，而是在动作空间中做有界修正", "02 / Architecture")
    y = 2.3
    boxes = [
        ("环境状态", "gripper / object /\ngoal / velocity", 0.6, TEAL_LIGHT, TEAL),
        ("Human Intent", "4-stage 状态机\n解析式 P 控制", 3.0, ORANGE_LIGHT, ORANGE),
        ("SAC Assist", "MultiInputPolicy\n输出 assist action", 5.45, TEAL_LIGHT, TEAL),
        ("动作融合", "clip(a_h + 0.15 a_RL)\n人类主导", 7.9, OFF_WHITE, NAVY_2),
        ("MuJoCo", "接触动力学\n新状态 / reward", 10.35, ORANGE_LIGHT, ORANGE),
    ]
    for title, sub, x, fill, accent_color in boxes:
        add_process_box(slide, title, sub, x, y, 2.0, 1.35, fill, accent_color)
    for x1, x2 in [(2.6, 3.0), (5.0, 5.45), (7.45, 7.9), (9.9, 10.35)]:
        add_connector(slide, x1, y + 0.68, x2, y + 0.68, NAVY_2, 1.8)
    add_connector(slide, 11.35, y + 1.38, 11.35, 5.25, MID_GRAY, 1.5, False)
    add_connector(slide, 11.35, 5.25, 1.6, 5.25, MID_GRAY, 1.5, False)
    add_connector(slide, 1.6, 5.25, 1.6, y + 1.38, MID_GRAY, 1.5)
    add_text(slide, "闭环反馈", 5.75, 5.02, 1.7, 0.28, 11, MID_GRAY, True, PP_ALIGN.CENTER)
    add_rect(slide, 1.05, 5.75, 11.2, 0.8, NAVY, NAVY, True)
    add_text(
        slide,
        "核心设计原则：a_full = clip(a_human + assist_scale × a_assist),  assist_scale = 0.15",
        1.28,
        5.96,
        10.75,
        0.36,
        17,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        FONT_MONO,
    )
    add_footer(slide, 3, "Source: assistive_fetch/wrappers.py:135–155; fetch_push_wrapper.py:162–170")

    # 4. Human controller
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "Human Intent Controller：可解释的四阶段状态机", "先构造可用但不完美的人类基线，再让 RL 学习补偿", "03 / Control")
    phases = [
        ("0  LIFT", "抬升到安全高度\nz → 0.50 m", TEAL),
        ("1  APPROACH", "移动到物体后方\np_pre = p_obj − d·ê_goal", ORANGE),
        ("2  DESCEND", "保持 XY 对齐\nz → 0.425 m", GREEN),
        ("3  PUSH", "比例推向目标\nu_xy = Kp·(p_goal−p_obj)", NAVY_2),
    ]
    for i, (title, body, accent_color) in enumerate(phases):
        x = 0.55 + i * 3.16
        add_rect(slide, x, 1.65, 2.72, 2.0, WHITE, accent_color, True)
        add_text(slide, title, x + 0.18, 1.88, 2.36, 0.34, 16, accent_color, True, PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.18, 2.4, 2.36, 0.9, 12.5, DARK, False, PP_ALIGN.CENTER)
        if i < 3:
            add_connector(slide, x + 2.76, 2.65, x + 3.10, 2.65, MID_GRAY, 1.8)
    add_text(slide, "关键控制细节", 0.65, 4.02, 2.8, 0.38, 19, NAVY, True)
    add_rich_lines(
        slide,
        [
            "behind_offset = 0.08 m：从目标反方向构造预推位置，确保从物体后方接触。",
            "阈值与滞回：XY、Z 达到阈值后切换阶段；严重失配时从 PUSH 回退到 APPROACH。",
            "Debug 发现：behind_offset 与回退阈值冲突曾导致“刚进 PUSH 就退出”；修正后各阶段进入率接近 1.0。",
            "当前局限：垂直下降可能碰到立方体边缘，造成弹起；后续应加入接触感知与连续过渡。",
        ],
        0.72,
        4.55,
        11.85,
        1.95,
        13,
    )
    add_footer(slide, 4, "Source: assistive_fetch/fetch_push_wrapper.py:64–160")

    # 5. RL formulation
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "RL 问题建模：让策略只学习“修正量”", "降低搜索空间、保留可解释人类意图、约束辅助权限", "04 / RL Formulation")
    add_card(slide, "观测 sₜ", "Fetch 原始状态\n+ 当前 human action\n+ 上一步 assist action\n+ goal error", 0.65, 1.55, 3.55, 2.65, TEAL)
    add_card(slide, "策略输出 aᴿᴸ", "4D 连续动作\n[x, y, z, gripper]\n经 tanh / clip 限幅\n再乘 assist_scale=0.15", 4.55, 1.55, 3.55, 2.65, ORANGE)
    add_card(slide, "执行动作 a_full", "a_full = clip(a_h + 0.15aᴿᴸ)\n\n人类提供任务意图；\n策略学习纠偏、加速和接触补偿", 8.45, 1.55, 3.55, 2.65, GREEN)
    add_text(slide, "为什么这种分解适合运动控制？", 0.68, 4.62, 4.5, 0.38, 19, NAVY, True)
    add_rich_lines(
        slide,
        [
            "结构先验：human controller 提供方向和阶段，RL 不必从零发现完整操作策略。",
            "安全边界：辅助幅值可显式限制，便于分析最大干预量。",
            "可解释评估：可分别统计 human effort、assist effort、full action 与 jerk。",
        ],
        0.75,
        5.08,
        7.3,
        1.48,
        13,
    )
    add_rect(slide, 8.55, 4.72, 3.55, 1.65, NAVY, NAVY, True)
    add_text(slide, "SAC", 8.75, 4.95, 3.15, 0.46, 25, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Off-policy · 连续控制\n200k steps 内获得有效策略", 8.75, 5.52, 3.15, 0.62, 12, RGBColor(201, 220, 228), False, PP_ALIGN.CENTER)
    add_footer(slide, 5, "Source: assistive_fetch/wrappers.py:41–58, 135–155; train_assistive_fetch.py:65–78")

    # 6. Reward
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "奖励设计：任务性能、辅助成本与平滑性的权衡", "最终方案使用纯 dense SAC；移除 HER，避免自定义奖励重标不一致", "05 / Reward")
    add_rect(slide, 0.8, 1.55, 11.75, 0.95, NAVY, NAVY, True)
    add_text(
        slide,
        "rₜ = −w_d ‖p_goal − p_obj‖ + b_s·𝟙_success − w_a‖a_assist‖ − w_Δ‖a_assist,t − a_assist,t−1‖",
        1.05,
        1.83,
        11.25,
        0.42,
        17,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        FONT_MONO,
    )
    terms = [
        ("距离项", "w_d = 1.0\n每步提供方向明确的 dense 信号", TEAL),
        ("成功奖励", "b_s = 10\n强化进入 5 cm 成功区域", GREEN),
        ("辅助成本", "w_a = 0.002–0.005\n抑制无意义大幅辅助", ORANGE),
        ("平滑项", "w_Δ = 0.3\n惩罚相邻 assist 动作变化", RED),
    ]
    for i, (t, b, c) in enumerate(terms):
        x = 0.65 + i * 3.12
        add_card(slide, t, b, x, 2.95, 2.72, 1.65, c)
    add_text(slide, "关键迭代", 0.7, 5.02, 2.3, 0.35, 18, NAVY, True)
    add_rich_lines(
        slide,
        [
            "dist_progress 会发生轨迹内望远镜相消，单步信号小；改用绝对距离 −dist。",
            "自定义历史相关 reward 与 HER 重标语义不一致；最终使用普通 replay buffer。",
            "平滑系数提高后 jerk 下降，但成功率出现权衡；说明 reward 是多目标控制器而非单一指标。",
        ],
        0.75,
        5.42,
        11.75,
        1.08,
        12.5,
    )
    add_footer(slide, 6, "Source: fetch_push_wrapper.py:174–200; train_assistive_fetch.py:65–78")

    # 7. Debug evolution
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "工程 Debug：从“不收敛”到可解释的控制闭环", "通过 rollout 日志、阶段进入次数和 reward 分解定位问题", "06 / Debug")
    stages = [
        ("V1", "单阶段 human\n奖励不稳定", "现象：dist 无下降趋势", RED),
        ("V2", "两阶段 human\n相位逻辑失配", "发现：PUSH 立即回退", ORANGE),
        ("V3", "修正状态机\n绝对距离 reward", "结果：84% vs 12%", GREEN),
        ("V5", "加入 jerk 统计\n增大平滑正则", "结果：性能/平滑权衡", TEAL),
    ]
    for i, (v, title, note, c) in enumerate(stages):
        x = 0.58 + i * 3.17
        add_rect(slide, x, 1.75, 2.75, 2.12, WHITE, c, True)
        add_text(slide, v, x + 0.18, 1.95, 0.55, 0.38, 17, c, True)
        add_text(slide, title, x + 0.75, 1.92, 1.75, 0.62, 14, NAVY, True)
        add_text(slide, note, x + 0.18, 2.82, 2.35, 0.52, 11.5, DARK)
        if i < 3:
            add_connector(slide, x + 2.78, 2.8, x + 3.12, 2.8, MID_GRAY, 1.7)
    add_text(slide, "Debug 方法论", 0.68, 4.32, 2.8, 0.38, 19, NAVY, True)
    add_card(slide, "1. 可观测", "逐步输出 dist、phase、XY/Z error、human/assist/full action norm。", 0.65, 4.82, 3.75, 1.55, TEAL)
    add_card(slide, "2. 可归因", "按 success/failure 分组，比较阶段进入次数与 reward 分项。", 4.78, 4.82, 3.75, 1.55, ORANGE)
    add_card(slide, "3. 可复现", "相同 seed 对比 model+assist 与 human-only，保存 checkpoint 与独立 eval。", 8.9, 4.82, 3.75, 1.55, GREEN)
    add_footer(slide, 7, "Source: debug_rollout_assistive_fetch.py; eval_assistive_fetch.py")

    # 8. SAC vs PPO
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "算法选择：为什么当前阶段采用 SAC，而不是 PPO", "连续动作 + 有限交互预算下，样本效率优先", "07 / Algorithm")
    headers = ["维度", "SAC（本项目）", "PPO"]
    rows = [
        ["数据利用", "Off-policy，replay buffer 复用经验", "On-policy，数据使用后丢弃"],
        ["样本效率", "高；200k steps 获得有效策略", "通常需要更多交互样本"],
        ["连续控制", "熵正则，适合 MuJoCo 连续动作", "同样可用，但更新更保守"],
        ["时序扩展", "可用 frame stacking；SB3 无原生 recurrent SAC", "RecurrentPPO 可直接接 LSTM"],
        ["动作平滑", "算法本身不保证，需显式正则/滤波", "同样不保证，不是换算法即可解决"],
    ]
    x0, y0 = 0.65, 1.55
    widths = [2.0, 4.85, 4.85]
    row_h = 0.78
    for j, header in enumerate(headers):
        x = x0 + sum(widths[:j])
        add_rect(slide, x, y0, widths[j], row_h, NAVY, NAVY)
        add_text(slide, header, x + 0.08, y0 + 0.2, widths[j] - 0.16, 0.34, 14, WHITE, True, PP_ALIGN.CENTER)
    for i, row in enumerate(rows):
        y = y0 + row_h * (i + 1)
        fill = WHITE if i % 2 == 0 else OFF_WHITE
        for j, cell in enumerate(row):
            x = x0 + sum(widths[:j])
            add_rect(slide, x, y, widths[j], row_h, fill, LIGHT_GRAY)
            add_text(slide, cell, x + 0.1, y + 0.13, widths[j] - 0.2, 0.53, 11.5, NAVY if j == 0 else DARK, j == 0, PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT)
    add_rect(slide, 0.85, 6.5, 11.3, 0.55, TEAL_LIGHT, TEAL, True)
    add_text(slide, "结论：当前保留 SAC；域随机化阶段优先尝试 SAC + observation history，再评估 RecurrentPPO。", 1.05, 6.64, 10.9, 0.28, 13, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 8, "Configuration: SAC MultiInputPolicy · γ=0.98 · τ=0.01 · batch=256")

    # 9. Quantitative results
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "定量结果：辅助策略显著提高成功率并降低最终距离", "50 episodes 独立评估；human-only 使用零辅助动作", "08 / Results")
    add_native_chart(
        slide,
        ["V3 任务性能优先", "V5 平滑正则"],
        [("model+assist", [0.84, 0.70]), ("human-only", [0.12, 0.20])],
        0.55,
        1.45,
        5.95,
        3.25,
        "成功率对比",
        1.0,
        True,
    )
    add_native_chart(
        slide,
        ["V3 任务性能优先", "V5 平滑正则"],
        [("model+assist", [0.0350, 0.0528]), ("human-only", [0.1249, 0.1081])],
        6.83,
        1.45,
        5.95,
        3.25,
        "最终目标距离（m，越低越好）",
        0.15,
        False,
    )
    add_metric(slide, "+72 pp", "V3 成功率提升", 0.85, 5.05, 2.5, TEAL, "84% vs 12%")
    add_metric(slide, "−72%", "V3 最终距离下降", 3.7, 5.05, 2.5, GREEN, "0.0350 m vs 0.1249 m")
    add_metric(slide, "−18%", "V3 human effort 下降", 6.55, 5.05, 2.5, ORANGE, "0.1812 vs 0.2208")
    add_metric(slide, "70%", "V5 平滑正则成功率", 9.4, 5.05, 2.5, NAVY_2, "性能与平滑性仍需权衡")
    add_footer(slide, 9, "Source: eval_assistive_fetch.py · V3/V5 50-episode comparison")

    # 10. Smoothness
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "平滑性实验：显式 jerk 正则有效，但尚未解决执行抖动", "V5：smoothness_coef = 0.3", "09 / Smoothness")
    image_path = LOG_DIR / "v5_with_increased_smoothness.png"
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.55), Inches(1.45), width=Inches(7.55))
    add_rect(slide, 8.42, 1.45, 4.25, 4.58, WHITE, LIGHT_GRAY, True)
    add_text(slide, "独立评估", 8.7, 1.7, 3.7, 0.35, 18, NAVY, True)
    add_metric(slide, "0.0925", "Assist jerk", 8.65, 2.22, 1.7, TEAL, "‖Δassist‖ / step")
    add_metric(slide, "0.1254", "Full jerk", 10.55, 2.22, 1.7, RED, "assist-on")
    add_metric(slide, "0.0477", "Full jerk", 8.65, 3.55, 1.7, GREEN, "human-only")
    add_metric(slide, "70%", "Success", 10.55, 3.55, 1.7, NAVY_2, "assist-on")
    add_text(
        slide,
        "解释",
        8.7,
        4.95,
        1.2,
        0.3,
        15,
        NAVY,
        True,
    )
    add_text(
        slide,
        "解析式 human 控制天然连续；神经网络 assist 对接触和观测变化进行高频修正。当前只惩罚 Δassist，尚未直接优化 Δfull。",
        8.7,
        5.3,
        3.6,
        0.62,
        11.5,
        DARK,
    )
    add_rect(slide, 0.65, 6.25, 12.0, 0.62, ORANGE_LIGHT, ORANGE, True)
    add_text(slide, "下一步：将 full-action jerk 纳入目标，或在执行端加入可验证的低通滤波 / rate limiter。", 0.9, 6.43, 11.5, 0.28, 13, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 10, "Note: 图标题中的“SAC+HER”为旧版 callback 文案；当前训练代码为 SAC without HER")

    # 11. Next steps
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_title(slide, "下一步：从单一 Human 扩展到人群分布与鲁棒共享控制", "目标：对不同力量、偏置、噪声和反应特性的用户自适应", "10 / Roadmap")
    roadmap = [
        ("Human 域随机化", "gain / bias / noise / delay\n建立人群行为分布", TEAL),
        ("时序意图估计", "frame stacking 或 recurrent policy\n在线识别用户特征", ORANGE),
        ("控制品质约束", "full-action jerk / rate limit\n接触冲击与弹起统计", GREEN),
        ("鲁棒评估", "固定 seed + 分层人群测试\n成功率 / effort / jerk / safety", NAVY_2),
    ]
    for i, (t, b, c) in enumerate(roadmap):
        x = 0.58 + i * 3.17
        add_rect(slide, x, 1.7, 2.75, 1.75, WHITE, c, True)
        add_text(slide, f"{i + 1}", x + 0.16, 1.9, 0.42, 0.42, 17, c, True, PP_ALIGN.CENTER)
        add_text(slide, t, x + 0.62, 1.87, 1.9, 0.42, 15, NAVY, True)
        add_text(slide, b, x + 0.18, 2.46, 2.35, 0.65, 11.5, DARK, False, PP_ALIGN.CENTER)
        if i < 3:
            add_connector(slide, x + 2.78, 2.58, x + 3.12, 2.58, MID_GRAY, 1.7)
    add_text(slide, "面向真实运动控制的补强", 0.65, 4.0, 4.4, 0.38, 19, NAVY, True)
    add_card(slide, "接触安全", "监测物体 z 速度/离桌高度，量化“边缘碰撞导致弹起”；加入冲击惩罚或接触状态机。", 0.65, 4.5, 3.72, 1.75, RED)
    add_card(slide, "时域设计", "评估 50→75 steps 与 γ 的配套调整；同时压缩 lift/descend 前置时间，避免仅靠延长时域。", 4.8, 4.5, 3.72, 1.75, ORANGE)
    add_card(slide, "Sim-to-Real", "随机化质量、摩擦、控制延迟和传感噪声；部署前增加动作限幅、watchdog 与 fallback。", 8.95, 4.5, 3.72, 1.75, TEAL)
    add_footer(slide, 11, "Roadmap: domain randomization → temporal adaptation → safety-constrained deployment")

    # 12. Summary
    slide = prs.slides.add_slide(blank)
    set_background(slide, NAVY)
    add_text(slide, "项目总结", 0.75, 0.55, 4.2, 0.62, 28, WHITE, True)
    add_text(slide, "从控制结构到学习闭环，我完成了：", 0.78, 1.35, 5.6, 0.4, 16, RGBColor(189, 211, 222))
    summary_items = [
        ("控制器设计", "基于几何关系构造预推点，设计四阶段状态机与 P 控制律。"),
        ("共享控制建模", "将 human intent 与受限 assist action 融合，保留解释性和控制权限边界。"),
        ("RL 工程落地", "使用 SAC、dense reward、checkpoint / callback / 独立评估完成训练闭环。"),
        ("系统化 Debug", "用逐步 rollout、phase entries、reward breakdown 定位相位与奖励问题。"),
        ("量化验证", "最佳成功率 84% vs 12%；同时识别 jerk 与接触安全的下一阶段问题。"),
    ]
    for i, (t, b) in enumerate(summary_items):
        y = 1.95 + i * 0.9
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.82), Inches(y), Inches(0.4), Inches(0.4))
        circle.fill.solid()
        circle.fill.fore_color.rgb = TEAL
        circle.line.fill.background()
        add_text(slide, t, 1.42, y - 0.02, 2.0, 0.35, 14, WHITE, True)
        add_text(slide, b, 3.35, y - 0.02, 6.55, 0.45, 12.5, RGBColor(224, 232, 238))
    add_rect(slide, 10.25, 1.75, 2.2, 3.8, NAVY_2, RGBColor(69, 96, 121), True)
    add_text(slide, "岗位匹配", 10.55, 2.02, 1.6, 0.4, 18, RGBColor(99, 215, 203), True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "运动学几何\n状态机控制\n连续控制 RL\n接触问题分析\n实验与指标设计\n工程 Debug",
        10.55,
        2.68,
        1.6,
        2.25,
        14,
        WHITE,
        False,
        PP_ALIGN.CENTER,
    )
    add_text(slide, "Q & A", 0.82, 6.85, 2.0, 0.35, 19, RGBColor(99, 215, 203), True)
    add_text(slide, "代码与结果均可现场展开讨论", 9.1, 6.9, 3.3, 0.25, 10, RGBColor(181, 192, 204), False, PP_ALIGN.RIGHT)

    # Basic document metadata
    props = prs.core_properties
    props.title = "基于共享控制与深度强化学习的 FetchPush 辅助操作"
    props.subject = "运动控制算法岗位面试项目汇报"
    props.author = "候选人"
    props.keywords = "Shared Control, SAC, FetchPush, Motion Control, Reinforcement Learning"
    props.comments = "Generated from the RL_robot Assistive Fetch project. All text, shapes, and charts are editable."

    prs.save(OUTPUT)
    return OUTPUT, len(prs.slides)


if __name__ == "__main__":
    output, count = build_presentation()
    print(f"Generated {output} with {count} slides")
