"""Add V3/V5 paired rollout animations to the interview presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from generate_assistive_fetch_ppt import (
    DARK,
    GREEN,
    LIGHT_GRAY,
    MID_GRAY,
    NAVY,
    OFF_WHITE,
    ORANGE,
    RED,
    TEAL,
    WHITE,
    add_card,
    add_footer,
    add_rect,
    add_text,
    add_title,
    set_background,
)


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "assistive_fetch_interview_presentation.pptx"
OUTPUT = ROOT / "assistive_fetch_interview_presentation_with_visualizations.pptx"
MEDIA_DIR = ROOT / "logs" / "assistive_fetch" / "visualizations"


def add_play_link(slide, x, y, mp4_path):
    button = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(2.25),
        Inches(0.45),
    )
    button.fill.solid()
    button.fill.fore_color.rgb = NAVY
    button.line.color.rgb = NAVY
    button.text_frame.clear()
    p = button.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "▶ 点击打开高清 MP4"
    run.font.name = "Microsoft YaHei"
    run.font.size = Inches(0.13)
    run.font.bold = True
    run.font.color.rgb = WHITE
    button.click_action.hyperlink.address = str(mp4_path.resolve())


def add_visualization_slide(prs, version, model_dist, human_dist, page_label):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFF_WHITE)
    gif_path = MEDIA_DIR / f"{version}_assist_vs_human.gif"
    mp4_path = MEDIA_DIR / f"{version}_assist_vs_human.mp4"
    if not gif_path.exists() or not mp4_path.exists():
        raise FileNotFoundError(f"Missing visualization media for {version}")

    subtitle = (
        "相同初始条件（seed=2）· 左：model+assist · 右：human-only · 50-step episode"
    )
    add_title(
        slide,
        f"{version.upper()} Visualization：辅助策略与 Human-only 对比",
        subtitle,
        "Visualization",
    )
    slide.shapes.add_picture(
        str(gif_path),
        Inches(0.55),
        Inches(1.42),
        width=Inches(8.55),
        height=Inches(4.82),
    )
    add_text(
        slide,
        "动画在 PowerPoint 放映模式下自动播放",
        0.72,
        6.35,
        4.25,
        0.25,
        9,
        MID_GRAY,
    )
    add_play_link(slide, 5.65, 6.24, mp4_path)

    add_rect(slide, 9.42, 1.42, 3.32, 4.82, WHITE, LIGHT_GRAY, True)
    add_text(slide, "代表性回合结果", 9.7, 1.72, 2.75, 0.36, 17, NAVY, True)
    add_text(
        slide,
        f"{model_dist:.3f} m",
        9.65,
        2.35,
        1.15,
        0.48,
        23,
        TEAL,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "model+assist\n成功",
        9.62,
        2.87,
        1.25,
        0.62,
        11,
        GREEN,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        f"{human_dist:.3f} m",
        11.15,
        2.35,
        1.15,
        0.48,
        23,
        ORANGE,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "human-only\n失败",
        11.12,
        2.87,
        1.25,
        0.62,
        11,
        RED,
        True,
        PP_ALIGN.CENTER,
    )
    add_card(
        slide,
        "观察重点",
        "比较两侧末端轨迹、进入 PUSH 的时间，以及物体到红色目标点的最终距离。",
        9.68,
        3.75,
        2.8,
        1.25,
        TEAL,
    )
    add_card(
        slide,
        "说明",
        "该回合用于定性演示；总体性能结论仍以 50-episode paired evaluation 为准。",
        9.68,
        5.15,
        2.8,
        0.85,
        ORANGE,
    )
    add_text(
        slide,
        page_label,
        12.25,
        7.08,
        0.5,
        0.22,
        9,
        MID_GRAY,
        True,
        PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        f"Embedded GIF + external MP4: logs/assistive_fetch/visualizations/{version}_assist_vs_human.mp4",
        0.55,
        7.14,
        10.9,
        0.18,
        7.5,
        MID_GRAY,
    )
    return slide


def move_slide(prs, slide, new_index):
    slide_id = slide._element.getparent()
    slide_id.remove(slide._element)
    slide_id.insert(new_index, slide._element)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    v3_slide = add_visualization_slide(prs, "v3", 0.025, 0.202, "09A")
    v5_slide = add_visualization_slide(prs, "v5", 0.024, 0.202, "09B")

    # Insert the visualizations after the quantitative-results slide (index 8),
    # before the smoothness-analysis slide.
    sld_ids = prs.slides._sldIdLst
    v3_id = sld_ids[-2]
    v5_id = sld_ids[-1]
    sld_ids.remove(v3_id)
    sld_ids.remove(v5_id)
    sld_ids.insert(9, v3_id)
    sld_ids.insert(10, v5_id)

    prs.core_properties.comments = (
        "Includes editable diagrams/charts plus embedded animated GIF comparisons "
        "for V3 and V5. GIFs play in slideshow mode; buttons link to MP4 files."
    )
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
